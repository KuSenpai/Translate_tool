"""Write a finished chapter (Vietnamese blocks) to a new .docx file."""
from __future__ import annotations

import os
import tempfile
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt

from ..errors import DocxError
from ..logging_setup import get_logger
from .models import Block

log = get_logger("DOCX")

_ALIGN = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
          "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT}


def _add_runs(paragraph, block: Block) -> None:
    for r in block.runs:
        parts = r.text.split("\n")
        for k, part in enumerate(parts):
            run = paragraph.add_run(part)
            run.bold, run.italic, run.underline = r.b or None, r.i or None, r.u or None
            if k < len(parts) - 1:
                run.add_break(WD_BREAK.LINE)


def build_document(title: str, blocks: list[Block]) -> docx.Document:
    document = docx.Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)
    if title:
        document.add_heading(title, level=1)
    for b in blocks:
        if b.type == "heading":
            p = document.add_heading("", level=max(1, min(3, (b.level or 1) + 1)))
        else:
            p = document.add_paragraph()
        _add_runs(p, b)
        if b.align in _ALIGN:
            p.alignment = _ALIGN[b.align]
    return document


def write_chapter_docx(path: str | Path, title: str, blocks: list[Block], *, protect: list[Path] = ()) -> Path:
    """Atomically write to `path`. Refuses to overwrite any file in `protect` (the source book)."""
    path = Path(path).resolve()
    for p in protect:
        if Path(p).resolve() == path:
            raise DocxError("Từ chối ghi đè file Word gốc.", code="DOCX_PROTECTED")
    path.parent.mkdir(parents=True, exist_ok=True)
    document = build_document(title, blocks)
    fd, tmp = tempfile.mkstemp(suffix=".docx", dir=path.parent)
    os.close(fd)
    try:
        document.save(tmp)
        os.replace(tmp, path)
    except PermissionError:
        Path(tmp).unlink(missing_ok=True)
        raise DocxError(f"Không ghi được {path.name} — file có thể đang mở trong Word. Hãy đóng file rồi thử lại.",
                        code="DOCX_LOCKED")
    except OSError as e:
        Path(tmp).unlink(missing_ok=True)
        raise DocxError(f"Không ghi được file DOCX: {e}", code="DOCX_WRITE_FAILED")
    log.info("Exported %s (%d blocks)", path.name, len(blocks))
    return path
