import docx
import pytest

from app.document.chapter_parser import detect_lang, extract_chapter, match_heading, parse_chapters
from app.document.docx_reader import read_docx
from app.document.models import Block, Run
from app.errors import ChapterError, DocxError


def blk(text, heading=False, bold=False):
    return Block(type="heading" if heading else "paragraph", runs=[Run(text=text, b=bold)])


@pytest.mark.parametrize("text,num", [
    ("제 12 화", 12), ("제12화", 12), ("12화", 12), ("Chương 12", 12), ("CHƯƠNG 12", 12),
    ("Chapter 12", 12), ("Chương 12: Gặp gỡ", 12), ("제 12 화 - 새로운 시작", 12), ("[Chương 12]", 12),
    ("chuong 12", 12), ("Ch. 12", 12),
])
def test_heading_patterns(text, num):
    m = match_heading(blk(text))
    assert m and m.number == num


@pytest.mark.parametrize("text", [
    "Anh ấy đã 12 tuổi.", "12", "Chương 12 là chương dài nhất trong tập này, và mọi người đều mong chờ nó.",
    "Chương 12 đã kết thúc.", "그는 12화에서 죽었다.", "Dòng hai — chương 12 của đời anh.",
    "Chương 12 " + "x" * 200,
])
def test_not_heading(text):
    assert match_heading(blk(text)) is None


def test_loose_title_accepted_when_styled():
    assert match_heading(blk("Chương 12 Gặp gỡ định mệnh", heading=True)).title == "Gặp gỡ định mệnh"


def test_detect_lang():
    assert detect_lang("김수현은 창밖을 바라보았다") == "ko"
    assert detect_lang("Kim Soo-hyun nhìn ra ngoài") == "vi"
    assert detect_lang("... 12 ...") is None


def test_sample_book_structure(book):
    blocks = read_docx(book)
    r = parse_chapters(blocks)
    assert sorted(r.chapters) == [10, 11, 12, 13, 14, 15]
    for n in (10, 11, 12, 13, 15):
        assert r.chapters[n].confidence == "high", (n, r.chapters[n].warnings)
    assert r.chapters[14].confidence == "low"
    assert any("SUY LUẬN" in w for w in r.chapters[14].warnings)


def test_extract_chapter_12_preserves_format(book):
    blocks = read_docx(book)
    ch = extract_chapter(blocks, parse_chapters(blocks), 12)
    assert all("(12화)" in b.text for b in ch.ko_blocks)
    assert ch.vi_title == "Gặp gỡ"
    first = ch.vi_blocks[0].runs
    assert [(r.text, r.b, r.i) for r in first] == [("Mở đầu ", False, False), ("in đậm", True, False),
                                                    (", rồi ", False, False), ("in nghiêng", False, True), (".", False, False)]
    assert any("\n" in b.text for b in ch.vi_blocks)            # soft line break preserved
    assert ch.vi_blocks[-1].text == "Hết chương 12." and ch.vi_blocks[-1].runs[0].i
    assert not any("(13화)" in b.text or "Chương 13" in b.text for b in ch.vi_blocks)


def test_inferred_boundary(book):
    blocks = read_docx(book)
    ch = extract_chapter(blocks, parse_chapters(blocks), 14)
    assert all(detect_lang(b.text) == "ko" for b in ch.ko_blocks)
    assert ch.vi_blocks[0].text.startswith("Chương 14 là chương dài nhất")
    assert ch.confidence == "low"


def test_duplicates_flagged(dup_book):
    blocks = read_docx(dup_book)
    r = parse_chapters(blocks)
    assert r.chapters[1].confidence == "high"
    assert r.chapters[2].confidence == "low" and len(r.chapters[2].vi) == 2
    second = extract_chapter(blocks, r, 2, vi_choice=1)
    assert second.vi_blocks[0].text == "Bản dịch thứ hai của chương 2."


def test_missing_chapter(book):
    blocks = read_docx(book)
    with pytest.raises(ChapterError) as e:
        extract_chapter(blocks, parse_chapters(blocks), 99)
    assert e.value.code == "CHAPTER_NOT_FOUND"


def test_reader_errors(tmp_path):
    with pytest.raises(DocxError) as e:
        read_docx(tmp_path / "nope.docx")
    assert e.value.code == "DOCX_NOT_FOUND"
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(DocxError) as e:
        read_docx(bad)
    assert e.value.code == "DOCX_UNREADABLE"


def test_no_headings(tmp_path):
    d = docx.Document()
    d.add_paragraph("Chỉ có nội dung, không có heading.")
    p = tmp_path / "plain.docx"
    d.save(p)
    r = parse_chapters(read_docx(p))
    assert not r.chapters and r.warnings


def test_hyperlink_text_kept(tmp_path):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = docx.Document()
    d.add_paragraph("Chương 1")
    p = d.add_paragraph("Xem ")
    rid = p.part.relate_to("https://example.com", RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "liên kết"
    r.append(t)
    link.append(r)
    p._p.append(link)
    p.add_run(" này.")
    path = tmp_path / "link.docx"
    d.save(path)
    assert read_docx(path)[1].text == "Xem liên kết này."


@pytest.mark.parametrize("text,num,title", [
    ("1920. 다크 문", 1920, "다크 문"), ("1920. Trăng Rằm Đen", 1920, "Trăng Rằm Đen"),
    ("Tập 2244. Lezo", 2244, "Lezo"), ("EP.2221 2221. 경성 2033", 2221, "경성 2033"),
    ("Ads by Pubfuture 1921. 다크 문", 1921, "다크 문"), ("Interrupted EP.2281 2281. 신의 아틀란티스", 2281, "신의 아틀란티스"),
])
def test_numbered_title_headings(text, num, title):
    m = match_heading(blk(text))
    assert m and (m.number, m.title) == (num, title)


@pytest.mark.parametrize("text", [
    "100.000 chiến binh là quá nhiều. Cần giảm số lượng, và đây là cơ hội.",
    "17화. 최종 보스는커녕 중간보스가 누군지도 밝혀지지 않았다. 작품의 자세한 설정도 밝혀지지 않은 것이다. 죽은 작가를 되살리지 않는 한 알아낼 방도도 없었다.",
    "1. Đi chợ mua rau.",
])
def test_numbered_not_heading(text):
    assert match_heading(blk(text)) is None


def test_weak_heading_needs_sequence():
    body = lambda s: [blk(f"{s} {k}") for k in range(4)]
    blocks = ([blk("1920. 다크 문")] + body("가나다라마바사") + [blk("1920. Trăng Rằm Đen")] + body("Nội dung chương")
              + [blk("3. Tấn công")] + body("Tiếp tục nội dung")
              + [blk("1921. 다크 문")] + body("가나다라마바사") + [blk("1921. Trăng Rằm Đen")] + body("Nội dung chương"))
    r = parse_chapters(blocks)
    assert sorted(r.chapters) == [1920, 1921]
    assert any("3. Tấn công" in w for w in r.chapters[1920].warnings)          # reported, not silent
    ch = extract_chapter(blocks, r, 1920)
    assert any(b.text == "3. Tấn công" for b in ch.vi_blocks)                  # stays inside chapter 1920


def test_bare_number_heading_only_when_next_chapter():
    body = lambda s: [blk(f"{s} {k}") for k in range(4)]
    blocks = ([blk("2150. 이터널 에덴")] + body("가나다라마바사") + [blk("2150. Eternal Eden")] + body("Nội dung 2150")
              + [blk("2151")] + body("Nội dung 2151")
              + [blk("2152. 이터널 에덴")] + body("가나다라마바사") + [blk("2152. Eternal Eden")] + body("Nội dung")
              + [blk("7")] + body("Vẫn là 2152"))
    r = parse_chapters(blocks)
    assert sorted(r.chapters) == [2150, 2151, 2152]
    c2150 = extract_chapter(blocks, r, 2150)
    assert not any("2151" in b.text for b in c2150.vi_blocks)          # 2151 no longer merged into 2150
    c2151 = extract_chapter(blocks, r, 2151)
    assert c2151.vi_blocks[0].text == "Nội dung 2151 0" and not c2151.ko_blocks
    assert any("chỉ là con số" in w for w in c2151.warnings)
    assert extract_chapter(blocks, r, 2152).vi_blocks[-1].text == "Vẫn là 2152 3"   # "7" stays content
