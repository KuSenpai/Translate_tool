"""Whole-novel AI translation (mock translator: TRANSLATE_PROVIDER=mock, no API calls)."""
import time

import docx
import pytest

from app.ai.novel_translator import (MODELS, SCHEMA, AnthropicChapterTranslator, LLMRefused, normalize_vi_heading,
                                     read_txt, split_source)
from app.document.chapter_parser import parse_chapters
from app.document.docx_reader import read_docx

KO_TXT = """작가의 말: 즐겁게 읽어주세요.

1920. 다크 문
김수현은 창밖을 바라보았다.
“형, 오늘도 늦었어?”
누나는 대답하지 않았다.

1921. 다크 문
그는 칼을 뽑았다.
REFUSE 이 장면은 거절된다.
피가 흘렀다.

1922. 다크 문
새벽이 밝았다.
김수현은 웃었다.
모든 것이 끝났다.
"""


def wait(client, jid, timeout=30):
    end = time.time() + timeout
    while time.time() < end:
        j = client.get(f"/api/translate/jobs/{jid}").json()
        if j["status"] in ("done", "error", "paused") and not j["running"]:
            return j
        time.sleep(0.2)
    raise AssertionError(j)


def test_split_txt(tmp_path):
    p = tmp_path / "ko.txt"
    p.write_text(KO_TXT, encoding="utf-8")
    chapters, warnings = split_source(read_txt(p))
    assert [c.number for c in chapters] == [1920, 1921, 1922]
    assert all(c.lang == "ko" and not c.has_translation for c in chapters)


def test_txt_encodings(tmp_path):
    p = tmp_path / "cp949.txt"
    p.write_bytes("1. 제목\n안녕하세요\n".encode("cp949"))
    assert [b.text for b in read_txt(p)][:2] == ["1. 제목", "안녕하세요"]


def test_translate_txt_end_to_end(client, tmp_path):
    info = client.get("/api/translate/info").json()
    assert info["default_model"] == "claude-opus-5-5" and "18+" in info["default_style"]
    meta = client.post("/api/translate/analyze", files={"file": ("truyen.txt", KO_TXT.encode("utf-8"))}).json()
    assert meta["todo"] == 3 and meta["kind"] == "txt" and meta["estimate"]["claude-opus-5-5"]["usd"] > 0
    job = client.post("/api/translate/jobs", json={"sid": meta["sid"], "model": "claude-opus-5-5", "concurrency": 1}).json()
    j = wait(client, job["id"])
    assert j["status"] == "done" and j["counts"] == {"done": 2, "failed": 1}
    failed = next(c for c in j["chapters"] if c["status"] == "failed")
    assert failed["number"] == 1921 and "AI_REFUSED" in failed["error"]
    assert client.get(f"/api/translate/jobs/{job['id']}").json()["auto_terms"] == {"김수현": "Kim Soo-hyun"}

    out_txt = open(j["output"]["txt"], encoding="utf-8").read()
    # Korean chapter, then its translation right under it; the refused chapter stays Korean-only.
    assert out_txt.index("1920. 다크 문") < out_txt.index("1920. Tiêu đề dịch thử") < out_txt.index("1921. 다크 문")
    assert out_txt.index("1921. 다크 문") < out_txt.index("1922. 다크 문") < out_txt.index("1922. Tiêu đề dịch thử")
    assert "Bản dịch thử đoạn 1 (" in out_txt and "1921. Tiêu đề dịch thử" not in out_txt
    assert out_txt.startswith("작가의 말")

    # The bilingual output is what the editor reads: chapters pair up KO + VI.
    r = parse_chapters(read_docx(j["output"]["docx"]))
    assert len(r.chapters[1920].ko) == 1 and len(r.chapters[1920].vi) == 1
    assert len(r.chapters[1922].vi) == 1 and not r.chapters[1921].vi

    # Open it in the editor as a project.
    pid = client.post(f"/api/translate/jobs/{job['id']}/open").json()["project_id"]
    ch = client.post(f"/api/projects/{pid}/chapters/1920/load", json={}).json()
    assert ch["draft"][0]["runs"][0]["text"].startswith("Bản dịch thử đoạn 1")
    assert ch["confidence"] == "high"
    assert client.get(f"/api/translate/jobs/{job['id']}/download?fmt=docx").content[:2] == b"PK"


def test_translate_docx_skips_already_translated(client, book):
    """The sample book is bilingual except chapter 14 (no Vietnamese heading): nothing to redo but 14."""
    with open(book, "rb") as f:
        meta = client.post("/api/translate/analyze", files={"file": ("book.docx", f)}).json()
    assert meta["already_translated"] == 5
    todo = [c["number"] for c in meta["chapters"] if c["lang"] == "ko" and not c["has_translation"]]
    assert todo == [14]
    job = client.post("/api/translate/jobs", json={"sid": meta["sid"], "concurrency": 2}).json()
    j = wait(client, job["id"])
    assert j["counts"] == {"done": 1}
    d = docx.Document(j["output"]["docx"])
    texts = [p.text for p in d.paragraphs]
    assert texts.count("Chương 12 - Gặp gỡ") == 1               # existing translations untouched
    i14 = texts.index("14. Tiêu đề dịch thử") if "14. Tiêu đề dịch thử" in texts else texts.index("Chương 14")
    assert texts[i14 + 1].startswith("Bản dịch thử đoạn 1")             # chapter 14 now translated


def test_range_and_errors(client):
    meta = client.post("/api/translate/analyze", files={"file": ("t.txt", KO_TXT.encode())}).json()
    r = client.post("/api/translate/jobs", json={"sid": meta["sid"], "from_number": 5000})
    assert r.json()["error"]["code"] == "NOTHING_TO_DO"
    job = client.post("/api/translate/jobs", json={"sid": meta["sid"], "from_number": 1922, "to_number": 1922}).json()
    assert [c["number"] for c in job["chapters"]] == [1922]
    assert client.post("/api/translate/analyze", files={"file": ("x.pdf", b"x")}).json()["error"]["code"] == "DOCX_BAD_TYPE"
    assert client.post("/api/translate/analyze", files={"file": ("x.txt", "không có heading".encode())}).json()["error"]["code"] == "NO_CHAPTERS"


def test_normalize_heading():
    from app.ai.novel_translator import SourceChapter
    ko = SourceChapter(0, 2221, 0, 1, 2, "EP.2221 2221. 경성 2033", "ko")
    assert normalize_vi_heading("Tập 2221. Kinh Thành 2033", ko) == "Tập 2221. Kinh Thành 2033"
    assert normalize_vi_heading("Kinh Thành 2033", ko) == "2221. Kinh Thành 2033"
    assert normalize_vi_heading("", ko) == "Chương 2221"


def test_anthropic_request_shape(monkeypatch):
    """Model, structured output, effort, refusal fallbacks and the adult-content system prompt are sent."""
    from types import SimpleNamespace
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    tr = AnthropicChapterTranslator("claude-opus-5-5", "medium")
    seen = {}

    class FakeStream:
        def __init__(self, **kw):
            seen.update(kw)

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def get_final_message(self):
            text = '{"heading": "1920. Trăng Rằm Đen", "paragraphs": ["Một", "Hai"], "new_terms": []}'
            return SimpleNamespace(stop_reason=seen.get("_stop", "end_turn"), content=[SimpleNamespace(type="text", text=text)],
                                   usage=SimpleNamespace(input_tokens=10, output_tokens=20, cache_creation_input_tokens=0,
                                                         cache_read_input_tokens=5), model="claude-opus-5-5", stop_details=None)

    monkeypatch.setattr(tr.client.beta.messages, "stream", lambda **kw: FakeStream(**kw))
    r = tr.translate(heading="1920. 다크 문", paragraphs=["하나", "둘"], glossary_text="", prev_tail=[], style="STYLE")
    assert r.paragraphs == ["Một", "Hai"] and r.usage["cache_read_input_tokens"] == 5
    assert seen["model"] == "claude-opus-5-5" and seen["max_tokens"] == 64000
    assert seen["output_config"] == {"format": {"type": "json_schema", "schema": SCHEMA}, "effort": "medium"}
    assert seen["fallbacks"] == "default" and seen["betas"] == ["server-side-fallback-2026-07-01"]
    system = seen["system"][0]["text"]
    assert "18+" in system and "explicit sexual content" in system and "STYLE" in system
    assert "thinking" not in seen                                   # Opus 5.5: thinking is always adaptive

    def refused(**kw):
        kw["_stop"] = "refusal"
        return FakeStream(**kw)
    monkeypatch.setattr(tr.client.beta.messages, "stream", refused)
    with pytest.raises(LLMRefused):
        tr.translate(heading="h", paragraphs=["x"], glossary_text="", prev_tail=[], style="")
    assert MODELS["claude-haiku-4-5"]["effort"] is False             # effort is rejected on Haiku 4.5
