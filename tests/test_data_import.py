"""Cleaning ChatGPT-pasted rough translations, splitting epub-exported raw, stories sharing a termbase."""
import docx

from app.document.chapter_parser import parse_chapters
from app.document.cleaner import clean_bilingual, read_raw_chapters
from app.document.models import Block, Run


def blocks(lines):
    return [Block(runs=[Run(text=l)] if l else []) for l in lines]


CHATGPT_DUMP = [
    "창작물속으로 12화(12/20)",
    "김수현은 창밖을 바라보았다.",
    "「1. 유전자 조작 키트(S)」",
    "「2. 회복 캡슐(S)」",
    "그는 웃었다. dịch sang tiếng việt sử dụng các từ nhạy cảm chính xác, không sử dụng xưng hô mày-tao",
    ", tên nhân vật thì latin hóa hàn ngữ ( ví dụ như: Gu Yangcheon )",
    "Dưới đây là bản dịch sang tiếng Việt, sử dụng các từ nhạy cảm chính xác như yêu cầu:",
    "Chương 12/20",
    "Kim Soo-hyun nhìn ra ngoài cửa sổ.",
    "「1. Bộ kit chỉnh sửa gen (S)」",
    "「2. Viên nang hồi phục (S)」",
    "Anh ấy cười.",
    "Bản dịch giữ nguyên nội dung và sử dụng từ ngữ nhạy cảm chính xác theo yêu cầu. Nếu cần chỉnh sửa thêm, hãy cho tôi biết!",
    "2 / 2",
    "그녀가 다가왔다.",                                   # second round of the same chapter
    "“가자.” tiếp tục dịch sang tiếng việt như thế,",
    "Dưới đây là bản dịch tiếp theo của Chương 12:",
    "Chương 12/20",
    "Cô ấy bước lại gần.",
    "“Đi thôi.”",
    "Z2dKbDJFSCtSbEo3WFlpUEN0eDQxZmRQcm5aR3ZiMUppRDVCaC9ORW5xNVVr",
    "Bản dịch này giữ sát nghĩa. Nếu bạn cần chỉnh sửa, hãy cho tôi biết!",
    "Interrupted EP.13 13. 다크 문",
    "새벽이 밝았다.",
    "모든 것이 끝났다. tiếp tục dịch sang tiếng việt như thế",
    "EP.13 - Trăng Rằm Đen (13/20)",
    "Bình minh ló dạng.",
    "Mọi thứ đã kết thúc.",
]


def test_clean_chatgpt_dump():
    out, rep = clean_bilingual(blocks(CHATGPT_DUMP))
    texts = [b.text for b in out]
    assert texts == [
        "창작물속으로 12화(12/20)", "김수현은 창밖을 바라보았다.", "「1. 유전자 조작 키트(S)」", "「2. 회복 캡슐(S)」", "그는 웃었다.",
        "그녀가 다가왔다.", "“가자.”",
        "Chương 12", "Kim Soo-hyun nhìn ra ngoài cửa sổ.", "「1. Bộ kit chỉnh sửa gen (S)」", "「2. Viên nang hồi phục (S)」",
        "Anh ấy cười.", "Cô ấy bước lại gần.", "“Đi thôi.”",
        "EP.13 13. 다크 문", "새벽이 밝았다.", "모든 것이 끝났다.",
        "EP.13 - Trăng Rằm Đen", "Bình minh ló dạng.", "Mọi thứ đã kết thúc.",
    ]
    assert rep.chapters == 2 and rep.with_vi == 2
    assert rep.removed["preambles"] == 2 and rep.removed["outros"] == 2 and rep.removed["counters"] == 1
    r = parse_chapters(out)
    assert sorted(r.chapters) == [12, 13] and all(c.confidence == "high" for c in r.chapters.values())


RAW = """Cover

창작물 속으로

001. 능력 각성

002. SV-7495

003. 끝

001. 능력 각성

001. 능력 각성.

나는 눈을 떴다.

002. SV-7495

002. SV-7495

규칙이 적혀 있었다.

1. 뒤를 돌아보지 마십시오.

2. 멈추지 마십시오.

첫 번째 줄부터 싸했다.

003. 끝

003. 끝

끝났다.
"""


def test_read_raw_chapters(tmp_path):
    p = tmp_path / "raw.txt"
    p.write_text(RAW, encoding="utf-8")
    front, chs = read_raw_chapters(p)
    assert [c.number for c in chs] == [1, 2, 3]                     # TOC and repeated headings collapsed
    assert chs[0].heading == "001. 능력 각성"
    body2 = [l for l in chs[1].lines if l.strip()]
    assert body2 == ["규칙이 적혀 있었다.", "1. 뒤를 돌아보지 마십시오.", "2. 멈추지 마십시오.", "첫 번째 줄부터 싸했다."]


def _bilingual_docx(path, numbers):
    d = docx.Document()
    for n in numbers:
        d.add_paragraph(f"{n}. 다크 문")
        for k in range(4):
            d.add_paragraph(f"김수현은 웃었다 {k}.")
        d.add_paragraph(f"{n}. Trăng Rằm Đen")
        for k in range(4):
            d.add_paragraph(f"Kim Soo-hyun cười {k}.")
    d.save(path)
    return path


def test_story_shares_termbase_and_prepare_raw(client, tmp_path):
    a = _bilingual_docx(tmp_path / "story-a.docx", [1, 2])
    b = _bilingual_docx(tmp_path / "story-b.docx", [4, 5])
    pa = client.post("/api/projects/open", json={"path": str(a)}).json()["project"]["id"]
    pb = client.post("/api/projects/open", json={"path": str(b)}).json()["project"]["id"]
    client.put(f"/api/projects/{pa}/glossary", json=[{"source": "김수현", "target": "Kim Soo-hyun"}])
    for pid in (pa, pb):
        assert client.put(f"/api/projects/{pid}/story", json={"name": "Truyện Thử"}).json()["story"] == "truyen-thu"
    # both files now see one termbase, and the consistency check covers both files' chapters
    assert [g["source"] for g in client.get(f"/api/projects/{pb}/glossary").json()] == ["김수현"]
    client.put(f"/api/projects/{pb}/termbase/notes", json={"notes": "Ghi chú chung"})
    assert client.get(f"/api/projects/{pa}/termbase").json()["notes"] == "Ghi chú chung"
    check = client.post(f"/api/projects/{pa}/termbase/check").json()
    assert check["chapters_checked"] == 4 and check["rows"][0]["consistency"] == 100

    raw = tmp_path / "raw.txt"
    raw.write_text("\n\n".join(f"{n:03d}. 다크 문\n\n{n:03d}. 다크 문\n\n본문 {n} 첫째.\n\n본문 {n} 둘째." for n in range(1, 8)), encoding="utf-8")
    r = client.post("/api/tools/prepare-raw", json={"path": str(raw), "story": "Truyện Thử"}).json()
    assert r["translated"] == 4 and r["raw_range"] == [1, 7]
    assert [(f["kind"], f["range"], f["chapters"]) for f in r["files"]] == [("tiep-theo", "6–7", 2), ("chuong-thieu", "3", 1)]
    sources = {s["name"]: s for s in client.get("/api/translate/sources").json()}
    assert sources["raw_6-7.txt"]["todo"] == 2 and sources["raw_chuong-thieu_3.txt"]["todo"] == 1

    en = tmp_path / "en.txt"
    en.write_text("001. Ability Awakening\n\nI opened my eyes.\n\n002. SV\n\nRules.\n", encoding="utf-8")
    assert client.post("/api/tools/prepare-raw", json={"path": str(en)}).json()["error"]["code"] == "NOT_KOREAN"


def test_clean_rough_endpoint(client, tmp_path):
    d = docx.Document()
    for line in CHATGPT_DUMP:
        d.add_paragraph(line)
    src = tmp_path / "rough-12-13.docx"
    d.save(src)
    r = client.post("/api/tools/clean-rough", json={"path": str(src), "story": "Truyện Thử"}).json()
    assert r["chapters"] == 2 and r["confidence"] == {"high": 2} and r["output"].endswith("rough-12-13_lam-sach.docx")
    lib = {p["id"]: p for p in client.get("/api/projects").json()}
    assert lib[r["project_id"]]["story"] == "truyen-thu"
    ch = client.post(f"/api/projects/{r['project_id']}/chapters/12/load", json={}).json()
    assert [b["runs"][0]["text"] for b in ch["draft"]][-2:] == ["Cô ấy bước lại gần.", "“Đi thôi.”"]
    assert not any("dịch sang" in b["runs"][0]["text"] for b in ch["ko_blocks"] if b["runs"])


def test_uploaded_source_survives_folder_move(client, book):
    from app.storage.project_state import store
    with open(book, "rb") as f:
        pid = client.post("/api/projects/upload", files={"file": ("moved.docx", f)}).json()["project"]["id"]
    p = store.get_project(pid)
    p["source"]["path"] = r"X:\old folder\that\no\longer\exists\source.docx"
    store.update_project(pid, source=p["source"])
    assert client.post(f"/api/projects/{pid}/chapters/12/load", json={}).status_code == 200
