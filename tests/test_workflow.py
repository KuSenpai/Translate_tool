"""Acceptance test: book.docx → upload → detect → chapter 12 → edit → save → export Chapter_12.docx."""
import docx


def err(resp):
    return resp.json()["error"]["code"]


def test_full_p0_workflow(client, book):
    original_bytes = book.read_bytes()
    with open(book, "rb") as f:
        r = client.post("/api/projects/upload", files={"file": ("wf-book.docx", f)})
    assert r.status_code == 200, r.text
    idx = r.json()
    pid = idx["project"]["id"]
    assert [c["number"] for c in idx["chapters"]] == [10, 11, 12, 13, 14, 15]

    r = client.post(f"/api/projects/{pid}/chapters/12/load", json={})
    ch = r.json()
    assert ch["status"] == "LOADED" and ch["title"] == "Chương 12: Gặp gỡ"
    assert "(12화)" in ch["ko_blocks"][0]["runs"][0]["text"]

    # Edit: fix a sentence, keep bold run, change title.
    blocks = ch["draft"]
    blocks[1]["runs"][0]["text"] = "Kim Soo-hyun nhìn ra ngoài khung cửa sổ."
    r = client.put(f"/api/projects/{pid}/chapters/12/draft", json={"blocks": blocks, "title": "Chương 12: Tái ngộ"})
    assert r.status_code == 200 and r.json()["status"] == "EDITING"

    # Draft persisted.
    assert client.get(f"/api/projects/{pid}/chapters/12").json()["draft"][1]["runs"][0]["text"].endswith("khung cửa sổ.")
    # Re-loading the chapter keeps the draft (Word source unchanged).
    assert client.post(f"/api/projects/{pid}/chapters/12/load", json={}).json()["draft"][1]["runs"][0]["text"].endswith("khung cửa sổ.")

    r = client.post(f"/api/projects/{pid}/chapters/12/export")
    assert r.status_code == 200, r.text
    out = r.json()["path"]
    assert out.endswith("Chapter_12.docx")
    d = docx.Document(out)
    texts = [p.text for p in d.paragraphs]
    assert texts[0] == "Chương 12: Tái ngộ"
    assert "Kim Soo-hyun nhìn ra ngoài khung cửa sổ." in texts
    bold = [r.text for p in d.paragraphs for r in p.runs if r.bold]
    italic = [r.text for p in d.paragraphs for r in p.runs if r.italic]
    assert "in đậm" in bold and "in nghiêng" in italic
    assert not any("김수현" in t for t in texts)                 # no Korean leaked into export
    assert book.read_bytes() == original_bytes                   # source untouched

    r = client.get(f"/api/projects/{pid}/chapters/12/docx")
    assert r.status_code == 200 and r.content[:2] == b"PK"


def test_state_machine(client, book):
    with open(book, "rb") as f:
        pid = client.post("/api/projects/upload", files={"file": ("sm-book.docx", f)}).json()["project"]["id"]
    client.post(f"/api/projects/{pid}/chapters/13/load", json={})
    url = f"/api/projects/{pid}/chapters/13/status"
    assert err(client.post(url, json={"event": "ready", "confirm": True})) == "INVALID_TRANSITION"
    assert client.post(url, json={"event": "review"}).json()["status"] == "REVIEWED"
    assert err(client.post(url, json={"event": "ready"})) == "CONFIRM_REQUIRED"
    assert client.post(url, json={"event": "ready", "confirm": True}).json()["status"] == "READY_TO_PUBLISH"
    # Editing after review sends the chapter back to EDITING.
    ch = client.get(f"/api/projects/{pid}/chapters/13").json()
    ch["draft"][0]["runs"][0]["text"] = "Sửa lại."
    assert client.put(f"/api/projects/{pid}/chapters/13/draft", json={"blocks": ch["draft"]}).json()["status"] == "EDITING"
    # Reset returns the Word original.
    r = client.post(f"/api/projects/{pid}/chapters/13/reset").json()
    assert r["draft"] == r["original_vi"]


def test_error_responses(client, tmp_path):
    r = client.post("/api/projects/open", json={"path": str(tmp_path / "missing.docx")})
    assert r.status_code == 422 and err(r) == "DOCX_NOT_FOUND"
    r = client.post("/api/projects/upload", files={"file": ("x.docx", b"garbage")})
    assert err(r) == "DOCX_UNREADABLE"
    r = client.post("/api/projects/upload", files={"file": ("x.txt", b"hello")})
    assert err(r) == "DOCX_BAD_TYPE"


def test_open_by_path_and_missing_chapter(client, book):
    pid = client.post("/api/projects/open", json={"path": str(book)}).json()["project"]["id"]
    r = client.post(f"/api/projects/{pid}/chapters/99/load", json={})
    assert r.status_code == 404 and err(r) == "CHAPTER_NOT_FOUND"


def test_glossary_validation(client, book):
    with open(book, "rb") as f:
        pid = client.post("/api/projects/upload", files={"file": ("gl-book.docx", f)}).json()["project"]["id"]
    g = [{"source": "김수현", "target": "Kim Soo-hyun", "avoid": ["Kim Su-hyeon"]},
         {"source": "선배", "target": "tiền bối"}]
    assert len(client.put(f"/api/projects/{pid}/glossary", json=g).json()) == 2
    ch = client.post(f"/api/projects/{pid}/chapters/10/load", json={}).json()
    blocks = ch["draft"]
    blocks[1]["runs"][0]["text"] = "Kim Su-hyeon nhìn ra ngoài cửa sổ."
    issues = client.post(f"/api/projects/{pid}/chapters/10/validate", json={"blocks": blocks}).json()
    assert any("Kim Su-hyeon" in i["message"] and i["block"] == 1 for i in issues)
    assert client.put(f"/api/projects/{pid}/glossary", json=[{"source": " ", "target": "x"}]).status_code == 422


def test_library_and_last_chapter(client, book):
    with open(book, "rb") as f:
        pid = client.post("/api/projects/upload", files={"file": ("lib-book.docx", f)}).json()["project"]["id"]
    client.post(f"/api/projects/{pid}/chapters/13/load", json={})
    client.post(f"/api/projects/{pid}/chapters/12/load", json={})
    client.post(f"/api/projects/{pid}/chapters/12/status", json={"event": "review"})
    lib = {p["id"]: p for p in client.get("/api/projects").json()}
    p = lib[pid]
    assert p["last_chapter"] == 12 and p["chapter_count"] == 6 and p["chapter_range"] == [10, 15]
    assert p["status_counts"] == {"LOADED": 1, "REVIEWED": 1} and p["source_exists"]
    assert client.get("/api/projects").json()[0]["id"] == pid          # most recently used first


def test_add_missing_chapter(client, book):
    with open(book, "rb") as f:
        pid = client.post("/api/projects/upload", files={"file": ("add-book.docx", f)}).json()["project"]["id"]
    add = lambda n, **kw: client.post(f"/api/projects/{pid}/chapters/{n}/add", json=kw)  # noqa: E731

    assert err(add(12, vi_text="x")) == "CHAPTER_EXISTS"                      # already in the Word file
    assert err(add(17)) == "EMPTY_CHAPTER"
    idx = add(17, ko_text="제 17 화\n\n안녕하세요.", vi_text="Xin chào.\nTạm biệt.", title="Chương 17: Thêm tay").json()
    assert idx["missing"] == [16] and any("Thiếu 1 " in w for w in idx["warnings"])
    c17 = next(c for c in idx["chapters"] if c["number"] == 17)
    assert c17["manual"] and c17["has_ko"] and c17["has_vi"]
    assert err(add(17, vi_text="y")) == "CHAPTER_EXISTS"                      # no silent overwrite

    ch = client.post(f"/api/projects/{pid}/chapters/17/load", json={}).json()
    assert ch["title"] == "Chương 17: Thêm tay" and len(ch["ko_blocks"]) == 2 and len(ch["draft"]) == 2
    assert client.put(f"/api/projects/{pid}/chapters/17/draft", json={"blocks": ch["draft"][:1]}).json()["status"] == "EDITING"
    assert len(client.post(f"/api/projects/{pid}/chapters/17/load", json={}).json()["draft"]) == 1   # edits survive a reload

    idx = add(16, vi_text="Chỉ có bản Việt.").json()
    assert idx["missing"] == [] and not any("Thiếu" in w for w in idx["warnings"])
    assert idx["project"]["chapter_range"] == [10, 17]
