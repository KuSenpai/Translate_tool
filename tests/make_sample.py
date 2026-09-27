"""Generate sample novel .docx files for tests and manual testing.

    python -m tests.make_sample            # writes samples/book.docx
"""
from __future__ import annotations

import sys
from pathlib import Path

import docx
from docx.enum.text import WD_BREAK

KO_BODY = [
    "김수현은 창밖을 바라보았다. 그는 열두 살 때부터 이 도시에 살았다.",
    "“형, 오늘도 늦었어?” 동생이 물었다.",
    "누나는 대답하지 않았다. 선배의 말이 머릿속에 맴돌았다.",
]
VI_BODY = [
    "Kim Soo-hyun nhìn ra ngoài cửa sổ. Anh đã sống ở thành phố này từ năm 12 tuổi.",
    "“Anh, hôm nay lại về muộn à?” Em trai hỏi.",
    "Chị không trả lời. Lời của tiền bối cứ văng vẳng trong đầu.",
]


def _para(document, text, bold=False, italic=False, style=None):
    p = document.add_paragraph(style=style)
    r = p.add_run(text)
    r.bold, r.italic = bold or None, italic or None
    return p


def build_book(path: Path) -> Path:
    d = docx.Document()
    _para(d, "Ghi chú của người dịch: file này chứa bản gốc và bản dịch.")
    ko_formats = ["제 {n} 화", "제{n}화", "{n}화", "제 {n} 화 - 새로운 시작", "제{n}화", "제 {n} 화"]
    vi_formats = ["Chương {n}", "CHƯƠNG {n}: Khởi đầu mới", "Chương {n} - Gặp gỡ", "Chapter {n}", "Chương {n}", "Chương {n}"]
    for k, n in enumerate(range(10, 16)):
        d.add_heading(ko_formats[k].format(n=n), level=2)
        for line in KO_BODY:
            _para(d, f"{line} ({n}화)")
        if n == 14:
            # No Vietnamese heading for chapter 14: the parser must infer (and warn).
            _para(d, f"Chương 14 là chương dài nhất trong tập này, và mọi người đều mong chờ nó.")
            for line in VI_BODY:
                _para(d, line)
            continue
        vi_heading = vi_formats[k].format(n=n)
        if n == 11:
            _para(d, vi_heading, bold=True)          # heading as bold paragraph, not a Word heading
        else:
            d.add_heading(vi_heading, level=2)
        p = d.add_paragraph()
        p.add_run("Mở đầu ")
        p.add_run("in đậm").bold = True
        p.add_run(", rồi ")
        p.add_run("in nghiêng").italic = True
        p.add_run(".")
        for line in VI_BODY:
            _para(d, line)
        # Soft line break inside one paragraph + a number that must NOT be a heading.
        p = d.add_paragraph()
        r = p.add_run("Dòng một của bài thơ")
        r.add_break(WD_BREAK.LINE)
        p.add_run("dòng hai — chương 12 của đời anh.")
        _para(d, "12")
        _para(d, f"Hết chương {n}.", italic=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)
    return path


def build_duplicate_book(path: Path) -> Path:
    d = docx.Document()
    for n in (1, 2):
        d.add_heading(f"제 {n} 화", level=2)
        _para(d, KO_BODY[0])
        d.add_heading(f"Chương {n}", level=2)
        _para(d, VI_BODY[0])
    d.add_heading("Chương 2", level=2)          # duplicate Vietnamese chapter 2
    _para(d, "Bản dịch thứ hai của chương 2.")
    d.save(path)
    return path


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "samples" / "book.docx"
    print(build_book(out))
